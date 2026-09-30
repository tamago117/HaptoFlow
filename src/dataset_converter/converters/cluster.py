"""Cluster dataset converter."""
import os
import sys
import numpy as np
import pandas as pd
import librosa
import scipy.interpolate
from typing import Dict, List, Tuple, Any
from multiprocessing import current_process

try:
	import openpyxl
	OPENPYXL_AVAILABLE = True
except ImportError:
	OPENPYXL_AVAILABLE = False

from src.dataset_converter.converters.base import BaseConverter
from src.dataset_converter.waveform import estimate_sample_rate
from src.dataset_converter.accel_transform import normalize_accel_transform, select_accel_streams
from src.utils.data.hf_dataset import resolve_dataset_root
from src.utils.data.sensor_io import (
	AUDIO_EXT,
	AUDIO_PATTERN,
	SENSOR_TABLE_EXT,
	SENSOR_TABLE_PATTERN,
	read_sensor_frame,
	sibling_sensor_path,
)


class ClusterDatasetConverter(BaseConverter):
	source = "cluster"

	def __init__(self, cfg: Dict):
		cluster_cfg = cfg.get("cluster_dataset", {})
		
		# Derive input.dir / input.pattern from sensor_type
		sensor_type = cluster_cfg.get("sensor_type", None)
		base_dataset_dir = None
		
		# "dft321" or a list of axes, e.g. ["X", "Y", "Z", "Norm"]
		self.accel_transform = normalize_accel_transform(cluster_cfg.get("accel_transform"))

		if sensor_type:
			sensor_type = sensor_type.lower()
			if sensor_type not in ("audio", "accel"):
				raise ValueError(f"sensor_type must be 'audio' or 'accel', got '{sensor_type}'")
			
			# Use base_dir if given, otherwise infer it from input.dir
			base_dataset_dir = cluster_cfg.get("base_dir", None)
			input_cfg = cfg.get("input", {})
			
			if base_dataset_dir:
				if sensor_type == "audio":
					input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "audio")
					input_cfg["pattern"] = AUDIO_PATTERN
				elif sensor_type == "accel":
					input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "accel")
					input_cfg["pattern"] = SENSOR_TABLE_PATTERN
			else:
				input_dir = input_cfg.get("dir", "")
				if input_dir:
					# .../sensor_data/audio -> dataset root; otherwise input.dir is the root
					if "sensor_data" in input_dir:
						base_dataset_dir = os.path.dirname(os.path.dirname(input_dir))
					else:
						base_dataset_dir = input_dir
					
					if sensor_type == "audio":
						input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "audio")
						input_cfg["pattern"] = AUDIO_PATTERN
					elif sensor_type == "accel":
						input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "accel")
						input_cfg["pattern"] = SENSOR_TABLE_PATTERN
				else:
					# Neither is set: use the default dataset root
					base_dataset_dir = resolve_dataset_root(allow_download=False)
					if sensor_type == "audio":
						input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "audio")
						input_cfg["pattern"] = AUDIO_PATTERN
					elif sensor_type == "accel":
						input_cfg["dir"] = os.path.join(base_dataset_dir, "sensor_data", "accel")
						input_cfg["pattern"] = SENSOR_TABLE_PATTERN
			
			cfg["input"] = input_cfg
		
		# Default output.dir: {base_dir}_{sensor_type}_converted
		if base_dataset_dir:
			output_cfg = cfg.get("output", {})
			if output_cfg.get("dir") is None:
				suffix = f"_{sensor_type}" if sensor_type else ""
				output_cfg["dir"] = f"{base_dataset_dir}{suffix}_converted"
				cfg["output"] = output_cfg
		
		super().__init__(cfg)
		self.texture_list_path = cluster_cfg.get(
			"texture_list_path",
			os.path.join(resolve_dataset_root(allow_download=False), "texture_list.xlsx"),
		)
		self.texture_map = self._load_texture_list()
		
	def _load_texture_list(self) -> Dict[str, str]:
		"""
		Load texture list from Excel file using openpyxl.
		Exits the program if the file cannot be loaded (file is corrupted).
		"""
		if not OPENPYXL_AVAILABLE:
			print(f"Error: openpyxl is not installed. Please install it with: pip install openpyxl")
			sys.exit(1)
		
		try:
			# data_only reads cached formula values; keep_vba=False avoids VBA parsing issues
			wb = openpyxl.load_workbook(
				self.texture_list_path,
				data_only=True,
				read_only=True,
				keep_vba=False
			)
			
			ws = wb.active
			
			headers = []
			for cell in ws[1]:
				headers.append(str(cell.value).strip() if cell.value else "")
			
			id_col_idx = None
			name_col_idx = None
			
			for idx, header in enumerate(headers):
				h_lower = header.lower()
				if ("no" in h_lower or "id" in h_lower) and id_col_idx is None:
					id_col_idx = idx
				if ("name" in h_lower or "material" in h_lower or "texture" in h_lower) and name_col_idx is None:
					if id_col_idx != idx:
						name_col_idx = idx
			
			if id_col_idx is None or name_col_idx is None:
				wb.close()
				print(f"Error: Could not identify ID/Name columns in {self.texture_list_path}")
				print(f"  Headers found: {headers}")
				print(f"  Expected columns containing 'no'/'id' and 'name'/'material'/'texture'")
				sys.exit(1)
			
			mapping = {}
			for row in ws.iter_rows(min_row=2, values_only=True):
				if row[id_col_idx] is not None and row[name_col_idx] is not None:
					texture_id = str(row[id_col_idx]).strip()
					material_name = str(row[name_col_idx]).strip()
					if texture_id and material_name:
						mapping[texture_id] = material_name
			
			wb.close()
			# Avoid duplicate output from worker processes
			if current_process().name == 'MainProcess':
				print(f"Loaded {len(mapping)} textures from {self.texture_list_path}")
			return mapping
			
		except FileNotFoundError:
			print(f"Error: Texture list file not found: {self.texture_list_path}")
			sys.exit(1)
		except Exception as e:
			print(f"Error: Failed to load texture list from {self.texture_list_path}")
			print(f"  Error type: {type(e).__name__}")
			print(f"  Error message: {e}")
			print(f"  The Excel file may be corrupted. Please check the file and try again.")
			sys.exit(1)

	def read_input(self, file_path: str) -> List[Tuple[np.ndarray, np.ndarray, Dict[str, Any], str]]:
		# Filename format: (texture_id)_(direction)_(velocity)_(force)_(repeat count).ext
		basename = os.path.basename(file_path)
		name_no_ext = os.path.splitext(basename)[0]
		parts = name_no_ext.split('_')
		
		if len(parts) < 5:
			print(f"Warning: Filename {basename} does not match expected format.")
			return []
			
		texture_id = parts[0]
		
		material_name = self.texture_map.get(texture_id, f"texture_{texture_id}")
		
		if file_path.endswith(AUDIO_EXT):
			waveform, sr = librosa.load(file_path, sr=self.sample_rate)
			duration = len(waveform) / sr
			timestamps = np.linspace(0, duration, len(waveform), dtype=np.float32)
			sensor_type = "audio"
			streams_data = [(waveform, "")]
		elif file_path.endswith(SENSOR_TABLE_EXT):
			df = read_sensor_frame(file_path)
			# Expected columns: time, X, Y, Z
			if "time" in df.columns:
				timestamps = df["time"].to_numpy(dtype=np.float32)
			else:
				# Fall back to the first column as time
				timestamps = df.iloc[:, 0].to_numpy(dtype=np.float32)
			
			X = df["X"].to_numpy(dtype=np.float32)
			Y = df["Y"].to_numpy(dtype=np.float32)
			Z = df["Z"].to_numpy(dtype=np.float32)

			# Remove DC offset (e.g., gravity on Z axis) before any further processing
			X = X - float(np.mean(X))
			Y = Y - float(np.mean(Y))
			Z = Z - float(np.mean(Z))

			fs = None
			try:
				fs = float(estimate_sample_rate(timestamps))
			except Exception:
				fs = float(self.sample_rate)
			if fs <= 0:
				fs = float(self.sample_rate)
			streams_data = select_accel_streams(X, Y, Z, fs, self.accel_transform)
			# DFT321 may return fewer samples than the input
			N = min(len(timestamps), min(len(w) for w, _ in streams_data))
			if N <= 0:
				return []
			timestamps = timestamps[:N]
			streams_data = [(w[:N], suffix) for w, suffix in streams_data]

			sensor_type = "accel"
		else:
			return []
		
		parent_dir = os.path.dirname(file_path) # e.g. .../audio/0
		grandparent_dir = os.path.dirname(parent_dir) # e.g. .../audio
		root_data_dir = os.path.dirname(grandparent_dir) # e.g. .../sensor_data
		
		# Force: .../sensor_data/force/{texture_id}/{filename_base}.parquet
		force_path = sibling_sensor_path(root_data_dir, "force", texture_id, name_no_ext)
		
		force_arr = np.zeros_like(timestamps)
		if os.path.exists(force_path):
			try:
				df_force = read_sensor_frame(force_path)
				if "force" in df_force.columns and "time" in df_force.columns:
					t_force = df_force["time"].to_numpy()
					f_force = df_force["force"].to_numpy()
					
					# Interpolate onto sensor timestamps, clamping to the edge values out of range
					f_force_interp = scipy.interpolate.interp1d(
						t_force, f_force, kind='linear', bounds_error=False, 
						fill_value=(f_force[0], f_force[-1])
					)
					force_arr = f_force_interp(timestamps).astype(np.float32)
				else:
					print(f"Warning: Force file {force_path} does not contain 'time' or 'force' column.")
			except Exception as e:
				print(f"Warning: Error processing force file {force_path}: {e}")
		else:
			# Missing force file: force stays 0
			pass
		
		# Velocity from .../sensor_data/position/{texture_id}/{filename_base}.parquet
		pos_path = sibling_sensor_path(root_data_dir, "position", texture_id, name_no_ext)
		
		vx_arr = np.zeros_like(timestamps)
		vy_arr = np.zeros_like(timestamps)
		
		if os.path.exists(pos_path):
			try:
				df_pos = read_sensor_frame(pos_path)
				t_pos = df_pos["time"].to_numpy()
				x_pos = df_pos["X"].to_numpy()
				y_pos = df_pos["Y"].to_numpy()
				
				# Position is recorded in mm; convert to m so velocity is in m/s
				x_pos_m = x_pos / 1000.0
				y_pos_m = y_pos / 1000.0
				
				# np.gradient gives d/dn; divide by dt for d/dt (position is ~100 Hz)
				dt_pos = np.gradient(t_pos)
				dt_pos[dt_pos == 0] = 1e-6
				
				vx_pos = np.gradient(x_pos_m) / dt_pos
				vy_pos = np.gradient(y_pos_m) / dt_pos
				
				# Same clamped interpolation as for force
				
				fx = scipy.interpolate.interp1d(t_pos, vx_pos, kind='linear', bounds_error=False, fill_value=(vx_pos[0], vx_pos[-1]))
				fy = scipy.interpolate.interp1d(t_pos, vy_pos, kind='linear', bounds_error=False, fill_value=(vy_pos[0], vy_pos[-1]))
				
				vx_arr = fx(timestamps).astype(np.float32)
				vy_arr = fy(timestamps).astype(np.float32)
				
			except Exception as e:
				print(f"Warning: Error processing position file {pos_path}: {e}")
		else:
			# Missing position file: velocities stay 0
			pass
			
		result = []
		for wav_data, suffix in streams_data:
			meta = {
				"label": material_name,
				"texture_id": texture_id,
				"force": force_arr,
				"velocity_x": vx_arr,
				"velocity_y": vy_arr,
				"sensor_type": sensor_type,
			}
			if suffix:
				meta["axis"] = suffix.strip("_")
				
			result.append((wav_data, timestamps, meta, suffix))
			
		return result
		
	def get_dataset_info_extras(self) -> Dict[str, Any]:
		return {
			"columns": ["timestamp", "force", "velocity_x", "velocity_y", "amplitude", "synthetic"]
		}

